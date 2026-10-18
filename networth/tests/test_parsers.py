from __future__ import annotations

import json
from pathlib import Path

import pytest

from networth import security
from networth.errors import UnreadableFileError, UnsupportedFileError
from networth.parsers import parse_file
from networth.parsers.document import column_letter
from networth.scripts.make_samples import write_pdf

SINGLE_SAMPLES = ["sample_input.csv", "sample_input.txt", "sample_input.xlsx",
                  "sample_input.docx", "sample_input.pdf"]


def all_text(path: Path) -> str:
    doc = parse_file(path)
    parts = [line.text for line in doc.lines]
    parts += [cell for t in doc.tables for row in t.rows for cell in row]
    return "\n".join(parts)


@pytest.mark.parametrize("name", SINGLE_SAMPLES + ["sample_couple.csv", "sample_couple.xlsx"])
def test_clean_samples_parse_and_pass_security(samples: Path, name: str) -> None:
    doc = parse_file(samples / name)
    assert not doc.is_empty
    assert security.scan(doc) == []


@pytest.mark.parametrize("name", SINGLE_SAMPLES)
def test_single_samples_contain_the_same_items(samples: Path, name: str) -> None:
    text = all_text(samples / name)
    for label in ["Bank Account 1", "Car 1", "401k 1", "Student Loan 1", "Credit Card 1", "712"]:
        assert label in text


@pytest.mark.parametrize("name", ["sample_with_pii.txt", "sample_with_pii.csv", "sample_with_pii.xlsx"])
def test_pii_samples_fail_security(samples: Path, name: str) -> None:
    assert security.scan(parse_file(samples / name))


def test_xlsx_numbers_lose_trailing_zero(samples: Path) -> None:
    doc = parse_file(samples / "sample_input.xlsx")
    row = next(r for r in doc.tables[0].rows if r[1] == "Bank Account 1")
    assert row[2] == "30000"


def test_pdf_lines_have_page_locations(samples: Path) -> None:
    doc = parse_file(samples / "sample_input.pdf")
    assert doc.lines[0].text == "Age: 29"
    assert doc.lines[0].location == "Page 1, Line 1"


def test_json(tmp_path: Path) -> None:
    path = tmp_path / "nw.json"
    path.write_text(json.dumps({"age": 29, "assets": [{"label": "Bank Account 1", "value": 30000.0}]}))
    doc = parse_file(path)
    assert ("label: Bank Account 1", "Field assets[0].label") in [(l.text, l.location) for l in doc.lines]
    assert "value: 30000" in [l.text for l in doc.lines]


def test_markdown_is_text(tmp_path: Path) -> None:
    path = tmp_path / "nw.md"
    path.write_text("# My net worth\n\n- Bank Account 1: $30,000\n")
    assert parse_file(path).lines[2].text == "- Bank Account 1: $30,000"


def test_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "nw.pages"
    path.write_text("x")
    with pytest.raises(UnsupportedFileError, match=r"\.pdf"):
        parse_file(path)


def test_missing_file(tmp_path: Path) -> None:
    with pytest.raises(UnreadableFileError, match="not found"):
        parse_file(tmp_path / "nope.csv")


def test_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("  \n\n")
    with pytest.raises(UnreadableFileError, match="no readable content"):
        parse_file(path)


@pytest.mark.parametrize("name", ["bad.xlsx", "bad.docx", "bad.pdf"])
def test_corrupted_file(tmp_path: Path, name: str) -> None:
    path = tmp_path / name
    path.write_bytes(b"SSN 123-45-6789 not really a document")
    with pytest.raises(UnreadableFileError) as exc:
        parse_file(path)
    assert "123-45-6789" not in str(exc.value)


def test_scanned_pdf_without_text(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    write_pdf(path, [])
    with pytest.raises(UnreadableFileError, match="scanned"):
        parse_file(path)


def test_latin1_text(tmp_path: Path) -> None:
    path = tmp_path / "nw.txt"
    path.write_bytes("Café fund 1: $500\n".encode("latin-1"))
    assert parse_file(path).lines[0].text == "Café fund 1: $500"


def test_column_letter() -> None:
    assert [column_letter(i) for i in (0, 25, 26, 27)] == ["A", "Z", "AA", "AB"]


def test_xls_without_xlrd_explains(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import builtins
    real_import = builtins.__import__

    def no_xlrd(name, *args, **kwargs):
        if name == "xlrd":
            raise ImportError("no xlrd")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_xlrd)
    path = tmp_path / "old.xls"
    path.write_bytes(b"not really excel")
    with pytest.raises(UnreadableFileError, match="save it as .xlsx"):
        parse_file(path)


def test_docx_tables_are_read(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Age: 30")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Label", "Value"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Bank Account 1", "$5,000"
    path = tmp_path / "nw.docx"
    document.save(str(path))
    doc = parse_file(path)
    assert doc.lines[0].text == "Age: 30"
    assert doc.tables[0].rows[1] == ("Bank Account 1", "$5,000")
    assert doc.tables[0].cell_location(1, 1) == "Table 1, Cell B2"
