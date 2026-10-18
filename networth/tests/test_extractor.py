from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest

from networth import extractor
from conftest import table_doc, text_doc
from networth.errors import InvalidInputError
from networth.models import Assumptions, Item, Person
from networth.parsers import parse_file

SINGLE_SAMPLES = ["sample_input.csv", "sample_input.txt", "sample_input.xlsx",
                  "sample_input.docx", "sample_input.pdf", "sample_input.json"]


def profile_of(ex: extractor.Extracted, household: str = "single", people: list[Person] | None = None):
    return extractor.build_profile(
        ex, household_type=household, people=people or [Person(29, 712)], state="TX",
        monthly_expenses=ex.monthly_expenses, assumptions=Assumptions(),
    )


# --- All formats agree ------------------------------------------------------

@pytest.mark.parametrize("name", SINGLE_SAMPLES)
def test_every_format_extracts_the_same_data(samples: Path, name: str) -> None:
    ex = extractor.extract(parse_file(samples / name))
    assert ex.ages == {1: "29"}
    assert ex.credit_scores == {1: "712"}
    assert ex.state in {"TX", "Texas"}
    assert ex.assets == [
        Item("Bank Account 1", 30000, "cash"),
        Item("Car 1", 2000, "vehicles"),
        Item("401k 1", 45000, "retirement"),
    ]
    assert ex.debts == [
        Item("Student Loan 1", 18000, "student", 5.5),
        Item("Credit Card 1", 4200, "credit_card", 24.9),
    ]
    assert ex.net_worth_history == {dt.date(y, 12, 31): v for y, v in
                                    [(2021, 8000), (2022, 19000), (2023, 31000), (2024, 42000), (2025, 54800)]}
    assert ex.income_history == {2021: 48000, 2022: 52000, 2023: 58000, 2024: 63000, 2025: 67000}
    assert ex.monthly_expenses == 3200
    assert ex.notes == []


def test_sample_profile_totals(samples: Path) -> None:
    profile = profile_of(extractor.extract(parse_file(samples / "sample_input.csv")))
    assert profile.total_assets == 77000
    assert profile.total_debts == 22200
    assert profile.net_worth == 54800


# --- Couples ----------------------------------------------------------------

@pytest.mark.parametrize("name", ["sample_couple.csv", "sample_couple.xlsx"])
def test_couple(samples: Path, name: str) -> None:
    ex = extractor.extract(parse_file(samples / name))
    assert ex.household_type == "Couple"
    assert ex.ages == {1: "34", 2: "31"}
    assert ex.credit_scores == {1: "745", 2: "668"}
    owners = {i.label: i.owner for i in (*ex.assets, *ex.debts)}
    assert owners == {"Bank Account 1": "joint", "401k 1": "partner1", "401k 2": "partner2",
                      "House 1": "joint", "Mortgage 1": "joint", "Student Loan 1": "partner2"}
    profile = profile_of(ex, "couple", [Person(34, 745), Person(31, 668)])
    assert profile.net_worth == 129000


def test_couple_income_is_summed_per_year() -> None:
    doc = table_doc(
        ["Section", "Label", "Value", "Owner"],
        ["Income", "2025", "60000", "Partner 1"],
        ["Income", "2025", "45000", "Partner 2"],
    )
    assert extractor.extract(doc).income_history == {2025: 105000}


def test_couple_needs_two_people(samples: Path) -> None:
    ex = extractor.extract(parse_file(samples / "sample_couple.csv"))
    with pytest.raises(InvalidInputError, match="2"):
        profile_of(ex, "couple", [Person(34, 745)])


def test_text_owner_in_parentheses() -> None:
    ex = extractor.extract(text_doc("ASSETS", "401k 1: $61,000 (Partner 1)"))
    assert ex.assets == [Item("401k 1", 61000, "retirement", owner="partner1")]


def test_bad_owner_is_rejected() -> None:
    doc = table_doc(["Section", "Label", "Value", "Owner"], ["Asset", "Car 1", "2000", "Partner 3"])
    with pytest.raises(InvalidInputError, match="Partner 1, Partner 2, or Joint"):
        extractor.extract(doc)


# --- Time series (date column + one column per account) ----------------------

@pytest.mark.parametrize("name", ["sample_timeseries.csv", "sample_timeseries.xlsx"])
def test_time_series(samples: Path, name: str) -> None:
    ex = extractor.extract(parse_file(samples / name))
    assert ex.as_of == dt.date(2025, 6, 15)
    assert ex.assets == [
        Item("Savings Account", 6000, "cash"),
        Item("Brokerage", 3600, "investments"),
        Item("Roth IRA", 2000, "retirement"),
    ]
    assert ex.debts == []  # the credit card is paid off in the latest row
    assert len(ex.net_worth_history) == 6
    assert ex.net_worth_history[dt.date(2025, 1, 15)] == 6500
    assert ex.net_worth_history[dt.date(2025, 6, 15)] == 11600
    assert ex.notes == []


def test_time_series_debt_column_and_total_mismatch() -> None:
    doc = table_doc(
        ["Date", "Bank Account 1", "Car Loan 1", "Total"],
        ["2025-01-01", "1000", "500", "500"],
        ["2025-02-01", "1500", "400", "9999"],
    )
    ex = extractor.extract(doc)
    assert ex.debts == [Item("Car Loan 1", 400, "auto")]
    assert ex.net_worth_history[dt.date(2025, 2, 1)] == 9999
    assert any("2025-02-01 (+$8,899)" in n for n in ex.notes)


def test_time_series_without_total_column_sums_columns() -> None:
    doc = table_doc(["Date", "Bank Account 1", "Credit Card 1"],
                    ["2025-01-01", "1000", "300"], ["2025-02-01", "1200", "100"])
    ex = extractor.extract(doc)
    assert ex.net_worth_history == {dt.date(2025, 1, 1): 700, dt.date(2025, 2, 1): 1100}


def test_time_series_rows_are_sorted_and_latest_used() -> None:
    doc = table_doc(["Date", "Brokerage", "Total"],
                    ["2025-03-01", "300", "300"], ["2025-01-01", "100", "100"])
    ex = extractor.extract(doc)
    assert ex.as_of == dt.date(2025, 3, 1)
    assert ex.assets == [Item("Brokerage", 300, "investments")]


# --- Text details -----------------------------------------------------------

def test_markdown_headings_and_bullets() -> None:
    ex = extractor.extract(text_doc(
        "# Profile", "- Age: 40", "- Credit Score (FICO): 801", "- State: Columbus, Ohio",
        "## Liabilities", "- Car Loan 1: $9,000 at 7.2%",
        "## Assets", "* Brokerage 1: $120k",
    ))
    assert ex.ages == {1: "40"} and ex.credit_scores == {1: "801"}
    assert ex.debts == [Item("Car Loan 1", 9000, "auto", 7.2)]
    assert ex.assets == [Item("Brokerage 1", 120000, "investments")]


def test_no_headings_guesses_from_labels() -> None:
    ex = extractor.extract(text_doc("Bank Account 1: $500", "Credit Card 1: $200", "2024: $300"))
    assert [i.label for i in ex.assets] == ["Bank Account 1"]
    assert [i.label for i in ex.debts] == ["Credit Card 1"]
    assert ex.net_worth_history == {dt.date(2024, 12, 31): 300}


def test_expense_lines_are_summed() -> None:
    ex = extractor.extract(text_doc("EXPENSES", "Rent: $1,500", "Food: $600"))
    assert ex.monthly_expenses == 2100


def test_negative_debt_becomes_balance_and_negative_asset_is_rejected() -> None:
    assert extractor.extract(text_doc("DEBTS", "Credit Card 1: -$4,200")).debts[0].value == 4200
    with pytest.raises(InvalidInputError, match="negative"):
        extractor.extract(text_doc("ASSETS", "Bank Account 1: -$100"))


def test_unreadable_rows_become_notes() -> None:
    ex = extractor.extract(text_doc("ASSETS", "Bank Account 1: lots", "Car 1: $2,000"))
    assert [i.label for i in ex.assets] == ["Car 1"]
    assert any("Bank Account 1" in n for n in ex.notes)


def test_assumptions_section() -> None:
    ex = extractor.extract(table_doc(
        ["Section", "Label", "Value"],
        ["Assumptions", "Investment Return", "6"],
        ["Assumptions", "Inflation", "3.5%"],
        ["Assumptions", "Warp Speed", "9"],
        ["Profile", "Monthly Contribution", "1500"],
    ))
    assert ex.assumptions == {"investment_return": "6", "inflation": "3.5%", "monthly_contribution": "1500"}
    assert any("Warp Speed" in n for n in ex.notes)


# --- build_profile notes ----------------------------------------------------

def test_notes_for_missing_data() -> None:
    profile = profile_of(extractor.extract(text_doc("ASSETS", "Brokerage 1: $1,000")))
    joined = " ".join(profile.notes)
    for phrase in ["No debts", "No cash", "Monthly expenses not given", "No income history", "Fewer than 2"]:
        assert phrase in joined


def test_no_items_is_an_error() -> None:
    with pytest.raises(InvalidInputError, match="No assets or debts"):
        profile_of(extractor.extract(text_doc("Age: 29")))


def test_history_sorted_in_profile() -> None:
    ex = extractor.extract(text_doc("ASSETS", "Bank Account 1: $5", "HISTORY", "2025: $5", "2023: $1"))
    assert list(profile_of(ex).net_worth_history) == [dt.date(2023, 12, 31), dt.date(2025, 12, 31)]


def test_docx_text_and_table_combine(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    for line in ("Age: 30", "State: Ohio", "Credit Score: 700"):
        document.add_paragraph(line)
    table = document.add_table(rows=3, cols=2)
    for r, (label, value) in enumerate([("Label", "Value"), ("Bank Account 1", "$5,000"), ("Car Loan 1", "$2,000")]):
        table.rows[r].cells[0].text, table.rows[r].cells[1].text = label, value
    path = tmp_path / "nw.docx"
    document.save(str(path))
    ex = extractor.extract(parse_file(path))
    assert ex.ages == {1: "30"} and ex.state == "Ohio"
    assert ex.assets == [Item("Bank Account 1", 5000, "cash")]
    assert ex.debts == [Item("Car Loan 1", 2000, "auto")]


def test_non_financial_dated_sheet_is_ignored() -> None:
    from networth.parsers.document import ParsedDocument, Table
    money = Table(rows=(("Date", "Brokerage", "Total"), ("2025-01-01", "100", "100"), ("2025-06-01", "300", "300")),
                  location_prefix="Sheet 'Net Worth', ")
    weight = Table(rows=(("Date", "Weight LBS", "Feedback"), ("2025-01-01", "268", ""), ("2025-02-01", "263", "ok")),
                   location_prefix="Sheet 'Weight', ")
    ex = extractor.extract(ParsedDocument(source=Path("x.xlsx"), tables=(money, weight)))
    assert ex.assets == [Item("Brokerage", 300, "investments")]
    assert ex.net_worth_history == {dt.date(2025, 1, 1): 100, dt.date(2025, 6, 1): 300}
    assert any("Ignored Sheet 'Weight'" in n and "Weight LBS" in n for n in ex.notes)


def test_unlabeled_accounts_with_a_total_still_count() -> None:
    ex = extractor.extract(table_doc(["Date", "Account A", "Total"], ["2025-01-01", "50", "50"],
                                     ["2025-03-01", "80", "80"]))
    assert ex.assets == [Item("Account A", 80, "other")]
