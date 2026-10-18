"""Excel (.xlsx, .xls) and CSV."""

from __future__ import annotations

import csv
import datetime as dt
import io
from pathlib import Path
from typing import Any

from networth.errors import UnreadableFileError
from networth.parsers.document import ParsedDocument, Table
from networth.parsers.text_parser import read_text


def cell_text(value: Any) -> str:
    """Spreadsheet cell -> text. Whole-number floats lose the ".0"; dates become ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    return str(value).strip()


def _table(rows: list[list[Any]], prefix: str) -> Table:
    return Table(rows=tuple(tuple(cell_text(c) for c in row) for row in rows), location_prefix=prefix)


def parse_csv(path: Path) -> ParsedDocument:
    rows = list(csv.reader(io.StringIO(read_text(path))))
    return ParsedDocument(source=path, tables=(_table(rows, ""),))


def parse_xlsx(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = wb.worksheets
        tables = tuple(
            _table(
                [list(r) for r in ws.iter_rows(values_only=True)],
                f"Sheet '{ws.title}', " if len(sheets) > 1 else "",
            )
            for ws in sheets
        )
    finally:
        wb.close()
    return ParsedDocument(source=path, tables=tables)


def parse_xls(path: Path) -> ParsedDocument:
    """Old .xls format. Needs the optional xlrd package."""
    try:
        import xlrd  # noqa: F401
    except ImportError:
        raise UnreadableFileError(
            "Reading old .xls files needs the 'xlrd' package. "
            "Open the file in Excel and save it as .xlsx, then try again."
        ) from None
    import pandas as pd

    frames = pd.read_excel(path, sheet_name=None, header=None, engine="xlrd")
    tables = tuple(
        _table(
            frame.astype(object).where(frame.notna(), None).values.tolist(),
            f"Sheet '{name}', " if len(frames) > 1 else "",
        )
        for name, frame in frames.items()
    )
    return ParsedDocument(source=path, tables=tables)
