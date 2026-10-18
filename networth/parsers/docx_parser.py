"""Word (.docx): paragraphs as lines, tables as tables."""

from __future__ import annotations

from pathlib import Path

from networth.parsers.document import Line, ParsedDocument, Table


def parse_docx(path: Path) -> ParsedDocument:
    from docx import Document

    document = Document(str(path))
    lines = tuple(
        Line(p.text, f"Paragraph {i}")
        for i, p in enumerate(document.paragraphs, start=1)
    )
    tables = tuple(
        Table(
            rows=tuple(tuple(cell.text.strip() for cell in row.cells) for row in t.rows),
            location_prefix=f"Table {i}, ",
        )
        for i, t in enumerate(document.tables, start=1)
    )
    return ParsedDocument(source=path, lines=lines, tables=tables)
