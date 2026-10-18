"""Regenerate the binary sample files (xlsx, docx, pdf) from the text samples.

Run from the project root: python -m networth.scripts.make_samples

All data is made up. The *_with_pii samples contain FAKE personal information
on purpose, to test that the security check rejects it.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from pathlib import Path

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


def write_pdf(path: Path, lines: list[str]) -> None:
    """Write a minimal one-page, text-based PDF (Helvetica). No dependencies."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    text_ops = "".join(f"({esc(line)}) Tj T*\n" for line in lines)
    stream = f"BT /F1 11 Tf 14 TL 50 760 Td\n{text_ops}ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def write_docx(path: Path, lines: list[str]) -> None:
    from docx import Document

    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(str(path))


def write_xlsx(path: Path, rows: list[list[str]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Net Worth"
    for row in rows:
        ws.append([_number_or_text(c) for c in row])
    wb.save(path)


def _number_or_text(cell: str) -> int | float | str | dt.datetime | None:
    if cell == "":
        return None
    if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", cell):
        return dt.datetime.strptime(cell, "%m/%d/%Y")  # real Excel dates, like a user's tracker
    try:
        return int(cell)
    except ValueError:
        try:
            return float(cell)
        except ValueError:
            return cell


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def _read_lines(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8").splitlines()


def main() -> None:
    write_xlsx(SAMPLES / "sample_input.xlsx", _read_csv(SAMPLES / "sample_input.csv"))
    write_xlsx(SAMPLES / "sample_couple.xlsx", _read_csv(SAMPLES / "sample_couple.csv"))
    write_xlsx(SAMPLES / "sample_with_pii.xlsx", _read_csv(SAMPLES / "sample_with_pii.csv"))
    write_xlsx(SAMPLES / "sample_timeseries.xlsx", _read_csv(SAMPLES / "sample_timeseries.csv"))
    text = _read_lines(SAMPLES / "sample_input.txt")
    write_docx(SAMPLES / "sample_input.docx", text)
    write_pdf(SAMPLES / "sample_input.pdf", text)
    print(f"Samples written to {SAMPLES}")


if __name__ == "__main__":
    main()
