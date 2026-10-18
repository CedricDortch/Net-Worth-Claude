"""Build data/*.json from official sources. Never hand-edit those files; rerun this.

Usage:
    python -m networth.scripts.build_reference_data --download SOURCES_DIR
    python -m networth.scripts.build_reference_data SOURCES_DIR

Run from the project root.

SOURCES_DIR holds the downloaded raw files (keep it outside the project; the SCF
file alone is ~20 MB). --download fetches them first. The raw files are treated
as untrusted data: only parsed, never executed.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import io
import json
import re
import urllib.request
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from networth.constants import STATES

DATA = Path(__file__).resolve().parents[1] / "data"

SCF_URL = "https://www.federalreserve.gov/econres/files/scfp2025excel.zip"
CENSUS_URL = ("https://www2.census.gov/programs-surveys/demo/tables/wealth/2023/"
              "wealth-asset-ownership/state_wealth_tables_dy2024.xlsx")
BEA_URL = "https://apps.bea.gov/regional/zip/SARPP.zip"
FRED_URL = "https://fred.stlouisfed.org/graph/fredgraph.csv?id={series}"
EXPERIAN_URL = "https://www.experian.com/blogs/ask-experian/what-is-the-average-credit-score-in-the-u-s/"

AGE_BRACKETS = (
    ("under_35", "Under 35", 0, 34), ("35_44", "35–44", 35, 44), ("45_54", "45–54", 45, 54),
    ("55_64", "55–64", 55, 64), ("65_74", "65–74", 65, 74), ("75_plus", "75 and older", 75, None),
)
FINE_BRACKETS = (("under_25", "Under 25", 0, 24), ("25_29", "25–29", 25, 29), ("30_34", "30–34", 30, 34))
PERCENTILES = (1, 5, 10, 20, 25, 30, 40, 50, 60, 70, 75, 80, 90, 95, 99)


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def _get(url: str, dest: Path) -> None:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (net-worth-analyzer data build)"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        dest.write_bytes(resp.read())
    print(f"  downloaded {dest.name}")


def download(src: Path) -> None:
    src.mkdir(parents=True, exist_ok=True)
    _get(SCF_URL, src / "scfp2025excel.zip")
    _get(CENSUS_URL, src / "state_wealth_tables_dy2024.xlsx")
    _get(BEA_URL, src / "SARPP.zip")
    _get(EXPERIAN_URL, src / "experian_average_credit_score.html")
    fred = src / "fred"
    fred.mkdir(exist_ok=True)
    for code in ["US", *STATES]:
        series = "MEHOINUSA646N" if code == "US" else f"MEHOINUS{code}A646N"
        _get(FRED_URL.format(series=series), fred / f"{code}.csv")


# ---------------------------------------------------------------------------
# Age percentiles (Federal Reserve SCF 2025)
# ---------------------------------------------------------------------------

def weighted_percentile(values: np.ndarray, weights: np.ndarray, p: float) -> float:
    """Smallest value whose cumulative weight share reaches p%."""
    order = np.argsort(values)
    v, w = values[order], weights[order]
    cum = np.cumsum(w) / w.sum()
    return float(v[min(np.searchsorted(cum, p / 100), len(v) - 1)])


def build_age_percentiles(src: Path) -> dict:
    with zipfile.ZipFile(src / "scfp2025excel.zip") as z:
        name = next(n for n in z.namelist() if n.lower().endswith(".csv"))
        with z.open(name) as f:
            df = pd.read_csv(f, usecols=["YY1", "WGT", "AGE", "NETWORTH", "INCOME"])

    def bracket(key: str, label: str, lo: int, hi: int | None) -> dict:
        s = df[(df.AGE >= lo) & (df.AGE <= (hi if hi is not None else 200))]
        v, w = s.NETWORTH.to_numpy(float), s.WGT.to_numpy(float)
        return {
            "key": key, "label": label, "min_age": lo, "max_age": hi,
            "percentiles": {str(p): round(weighted_percentile(v, w, p)) for p in PERCENTILES},
            "mean": round(float(np.average(v, weights=w))),
            "median_income": round(weighted_percentile(s.INCOME.to_numpy(float), w, 50)),
            "sample_households": int(s.YY1.nunique()),
        }

    return {
        "source": {
            "name": "Federal Reserve Board, 2025 Survey of Consumer Finances, Summary Extract Public Data",
            "url": SCF_URL,
            "page": "https://www.federalreserve.gov/econres/scfindex.htm",
            "year": 2025,
        },
        "retrieved": dt.date.today().isoformat(),
        "dollars": "2025 dollars",
        "unit": "household (family) net worth, grouped by age of the household head",
        "method": "Weighted percentiles over all 5 implicates using WGT. Values are close to, "
                  "but not identical to, the Fed's published figures (the public file is "
                  "disclosure-limited).",
        "brackets": [bracket(*b) for b in AGE_BRACKETS],
        "fine_brackets": [bracket(*b) for b in FINE_BRACKETS],
        "all": bracket("all", "All households", 0, None),
    }


# ---------------------------------------------------------------------------
# State data (Census SIPP wealth, BEA RPP, Census CPS income via FRED)
# ---------------------------------------------------------------------------

def _num(cell: object) -> float | None:
    """Census cells use '(B)' / '(D)' for suppressed values."""
    return float(cell) if isinstance(cell, (int, float)) else None


def _census_rows(ws) -> dict[str, tuple]:
    names = {n.lower(): c for c, n in STATES.items()}
    out = {}
    for row in ws.iter_rows(values_only=True):
        label = str(row[0]).strip().lower() if row[0] else ""
        if label == "total":
            out.setdefault("US", row)
        elif label in names:
            out[names[label]] = row
    return out


# Table 4 layout (0-based columns): 2 = zero or negative; then three groups, each with a
# total column and sub-range columns; 13 = $500,000 or over.
_T4_GROUPS = (
    (3, (4, 5, 6), (5_000, 10_000, 25_000)),   # $1–$24,999
    (7, (8, 9), (50_000, 100_000)),            # $25,000–$99,999
    (10, (11, 12), (250_000, 500_000)),        # $100,000–$499,999
)


def _distribution(row: tuple) -> dict[str, float] | None:
    """Cumulative % of households below each net worth bound ('0' = zero or negative).

    Uses sub-ranges when published, else the group total (sub-cells are often suppressed).
    """
    zero = _num(row[2])
    if zero is None:
        return None
    cumulative = {0: zero}
    running = zero
    for total_col, sub_cols, bounds in _T4_GROUPS:
        subs = [_num(row[c]) for c in sub_cols]
        total = _num(row[total_col])
        if all(s is not None for s in subs):
            for share, bound in zip(subs, bounds):
                running += share
                cumulative[bound] = running
        elif total is not None:
            running += total
            cumulative[bounds[-1]] = running
        else:
            break
    return {str(b): round(min(v, 100.0), 1) for b, v in cumulative.items()}


def build_state_data(src: Path) -> dict:
    wb = load_workbook(src / "state_wealth_tables_dy2024.xlsx", read_only=True, data_only=True)
    t1, t3, t4, t5 = (_census_rows(wb[n]) for n in ("Table 1", "Table 3", "Table 4", "Table 5"))

    with zipfile.ZipFile(src / "SARPP.zip") as z:
        name = next(n for n in z.namelist() if n.startswith("SARPP_STATE") and n.endswith(".csv"))
        text = z.read(name).decode("latin-1")
    rpp: dict[str, float] = {}
    names = {n.lower(): c for c, n in STATES.items()}
    for row in csv.DictReader(io.StringIO(text)):
        geo = (row.get("GeoName") or "").strip().lower()
        if (row.get("LineCode") or "").strip() == "1" and (geo in names or geo == "united states"):
            rpp["US" if geo == "united states" else names[geo]] = float(row["2024"])

    income: dict[str, tuple[int, float]] = {}
    for code in ["US", *STATES]:
        rows = list(csv.reader((src / "fred" / f"{code}.csv").read_text().splitlines()))[1:]
        date, value = next((r for r in reversed(rows) if r[1] not in ("", ".")))
        income[code] = (int(date[:4]), float(value))

    def entry(code: str) -> dict:
        return {
            "name": STATES.get(code, "United States"),
            "median_net_worth": _num(t1[code][1]),
            "mean_net_worth": _num(t5[code][1]),
            "net_worth_distribution": _distribution(t4[code]),
            "households_thousands": _num(t3[code][1]),
            "median_household_income": income[code][1],
            "income_year": income[code][0],
            "rpp": rpp[code],
        }

    return {
        "sources": [
            {"name": "U.S. Census Bureau, State-Level Wealth, Asset Ownership, & Debt of Households Tables: 2024 "
                     "(Survey of Income and Program Participation, survey year 2025)",
             "url": CENSUS_URL,
             "page": "https://www.census.gov/data/tables/2024/demo/wealth/state-wealth-asset-ownership.html",
             "fields": ["median_net_worth (Table 1)", "mean_net_worth (Table 5)",
                        "net_worth_distribution (Table 4)", "households_thousands (Table 3)"],
             "dollars": "2024 dollars", "released": "2026-07-15"},
            {"name": "U.S. Bureau of Economic Analysis, Regional Price Parities by State (SARPP), all items, 2024",
             "url": BEA_URL, "fields": ["rpp (United States = 100)"], "released": "2026-02-19"},
            {"name": "U.S. Census Bureau, Median Household Income by State (CPS ASEC), via FRED, "
                     "series MEHOINUS<ST>A646N (current dollars)",
             "url": "https://fred.stlouisfed.org/series/MEHOINUSA646N",
             "fields": ["median_household_income", "income_year"]},
        ],
        "retrieved": dt.date.today().isoformat(),
        "notes": [
            "null = suppressed by the Census Bureau (too few observations): '(B)' or '(D)' in the source.",
            "net_worth_distribution maps a dollar bound to the cumulative % of households below it "
            "('0' = zero or negative). Bounds whose sub-ranges were suppressed are omitted.",
            "SIPP (Census) and SCF (Federal Reserve) measure wealth differently; compare a state only "
            "with the national figures in this file, not with age_percentiles.json.",
        ],
        "national": entry("US"),
        "states": {code: entry(code) for code in STATES},
    }


# ---------------------------------------------------------------------------
# Credit scores (Experian)
# ---------------------------------------------------------------------------

def build_credit_averages(src: Path) -> dict:
    page = (src / "experian_average_credit_score.html").read_text(encoding="utf-8", errors="replace")
    names = {n.lower(): c for c, n in STATES.items()}
    states: dict[str, dict[str, int]] = {}
    years: list[str] = []
    for table in re.findall(r"<table.*?</table>", page, re.S):
        rows = [[html.unescape(re.sub(r"<[^>]+>", "", c)).strip()
                 for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", r, re.S)]
                for r in re.findall(r"<tr.*?</tr>", table, re.S)]
        if not rows or rows[0][0].lower() != "state":
            continue
        years = [h for h in rows[0][1:] if re.fullmatch(r"20\d\d", h)]
        for r in rows[1:]:
            code = names.get(r[0].lower())
            if code:
                states[code] = {y: int(v) for y, v in zip(years, r[1:1 + len(years)])}
        break
    if len(states) != len(STATES):
        raise SystemExit(f"Experian table: found {len(states)} states, expected {len(STATES)}. Page layout changed?")

    return {
        "source": {"name": "Experian, \"What Is the Average Credit Score in the U.S.?\" (average FICO Score, "
                           "September of each year)", "url": EXPERIAN_URL, "published": "2026-03-30"},
        "retrieved": dt.date.today().isoformat(),
        "score_model": "FICO Score 8",
        "latest_year": max(years),
        # National figures are stated in the article text, not the table.
        "national": {"2024": 715, "2025": 713},
        "states": states,
        "tiers": [
            {"name": "Poor", "min": 300, "max": 579},
            {"name": "Fair", "min": 580, "max": 669},
            {"name": "Good", "min": 670, "max": 739},
            {"name": "Very Good", "min": 740, "max": 799},
            {"name": "Exceptional", "min": 800, "max": 850},
        ],
        "tiers_source": "https://www.myfico.com/credit-education/what-is-a-fico-score",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("sources", type=Path, help="directory with the downloaded raw files")
    parser.add_argument("--download", action="store_true", help="download the raw files first")
    args = parser.parse_args()
    if args.download:
        download(args.sources)

    DATA.mkdir(exist_ok=True)
    for filename, build in (("age_percentiles.json", build_age_percentiles),
                            ("state_data.json", build_state_data),
                            ("credit_averages.json", build_credit_averages)):
        (DATA / filename).write_text(json.dumps(build(args.sources), indent=2, ensure_ascii=False) + "\n")
        print(f"wrote data/{filename}")


if __name__ == "__main__":
    main()
